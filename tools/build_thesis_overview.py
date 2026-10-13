"""Create the editable 10-minute thesis talk and matching PDF without LaTeX.

Run from the repository root with .venv/Scripts/python tools/build_thesis_overview.py.
"""
from pathlib import Path
import json
import textwrap
import fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'presentations'
ASSETS = OUT / 'thesis_overview_assets'
ASSETS.mkdir(exist_ok=True)
W, H = 13.333333, 7.5
NAVY, TEAL, INK, MUTED, LIGHT = '102A43', '007F86', '243B53', '52697E', 'EDF4F7'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
pdf = canvas.Canvas(str(OUT / 'Thesis_overview_10min.pdf'), pagesize=(W*72, H*72))
pdf.setTitle('Towards Fine-Grained Surgical Scene Understanding — Oluwatosin Alabi')
fontroot = Path('C:/Windows/Fonts')
pdfmetrics.registerFont(TTFont('Talk', str(fontroot / 'arial.ttf')))
pdfmetrics.registerFont(TTFont('TalkBold', str(fontroot / 'arialbd.ttf')))
notes, provenance = [], []

def rect(x,y,w,h,color, rounded=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(color)
    s.line.fill.background()
    pdf.setFillColor(HexColor('#'+color)); pdf.rect(x*72,(H-y-h)*72,w*72,h*72,fill=1,stroke=0)

def lines_for(text, width, size, bold):
    result=[]; font='TalkBold' if bold else 'Talk'
    for paragraph in text.split('\n'):
        line=''
        for word in paragraph.split():
            trial=(line+' '+word).strip()
            if line and pdfmetrics.stringWidth(trial,font,size)>width*72-8:
                result.append(line);line=word
            else: line=trial
        result.append(line)
    return result

def text(x,y,w,h,content,size=22,color=INK,bold=False):
    wrapped=lines_for(content,w,size,bold)
    needed=len(wrapped)*size*1.18/72
    if needed>h+0.04: raise ValueError(f'Text overflows: {content[:65]} ({needed:.2f}>{h})')
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(wrapped):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold
        p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(0);p.line_spacing=1.18
    pdf.setFillColor(HexColor('#'+color));pdf.setFont('TalkBold' if bold else 'Talk',size)
    for i,line in enumerate(wrapped):pdf.drawString(x*72,(H-y)*72-size-i*size*1.18,line)

def figure(rel,x,y,w,h):
    source=ROOT/rel;dest=ASSETS/(source.stem+'.png')
    doc=fitz.open(source);doc[0].get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).save(dest);doc.close()
    im=Image.open(dest); iw,ih=im.size;scale=min(w/iw,h/ih);dw,dh=iw*scale,ih*scale
    px,py=x+(w-dw)/2,y+(h-dh)/2
    slide.shapes.add_picture(str(dest), Inches(px), Inches(py),width=Inches(dw),height=Inches(dh))
    pdf.drawImage(str(dest),px*72,(H-py-dh)*72,dw*72,dh*72)

def takeaway(content):
    rect(.55,6.35,12.23,.58,LIGHT)
    text(.75,6.46,11.85,.4,content,16,TEAL,True)

def start(title,kicker,seconds,script,sources):
    global slide
    slide=prs.slides.add_slide(prs.slide_layouts[6]);idx=len(prs.slides)
    rect(0,0,W,H,'FFFFFF');rect(0,0,.15,H,TEAL)
    text(.55,.25,12,.3,kicker.upper(),12,TEAL,True)
    text(.55,.77,12.2,.85,title,30,NAVY,True)
    rect(.55,7.08,12.23,.012,'D5E0E8')
    text(.55,7.15,10,.22,'Oluwatosin Alabi  |  King’s College London  |  Thesis overview',10,MUTED)
    text(11.7,7.14,1,.25,f'{idx:02d} / 10',11,MUTED)
    slide.notes_slide.notes_text_frame.text=f'Target time: {seconds} seconds\n\n{script}\n\nSources: '+ '; '.join(sources)
    notes.append((idx,title,seconds,script));provenance.append({'slide':idx,'title':title,'sources':sources})

def end():pdf.showPage()

def card(x,y,w,h,title,body):
    rect(x,y,w,h,LIGHT);text(x+.2,y+.17,w-.4,.55,title,23,TEAL,True)
    text(x+.2,y+.87,w-.4,h-1.0,body,21)

start('Towards fine-grained surgical scene understanding','PhD thesis · 10-minute overview',30,
      'My thesis asks how a system can understand which individual surgical instrument is doing what, and to which tissue. The full thesis title is Towards Fine-Grained Surgical Scene Understanding through Grounded Instrument–Tissue Interaction Modelling. The work moves from a review of surgical vision to instance-level datasets, grounded action models, and finally a two-stage vision–language system. The unifying contribution is spatial grounding of surgical interactions.', ['Thesis.tex','Abstract.tex'])
text(.6,1.85,7.2,1.7,'Grounded instrument–tissue\ninteraction modelling',36,NAVY,True)
text(.6,3.85,7.2,1.05,'Oluwatosin Alabi\nSchool of Biomedical Engineering & Imaging Sciences',20)
text(.6,5.05,6.3,.9,'Supervisors: Prof. Tom Vercauteren\nand Prof. Miaojing Shi',19,MUTED)
card(8.1,2.0,4.6,3.8,'The central question','Which instrument instance\nperforms which action\non which anatomical target?')
takeaway('Review → instrument instances → grounded triplets → multimodal interaction prediction');end()

start('Recognising an action does not identify its actor','Problem · instrument, verb, target',60,
      'A surgical action triplet consists of an instrument, a verb, and a target: for example, a grasper retracting the liver. Traditional recognition tells us that this interaction occurs somewhere in a frame. But if two graspers appear, it does not identify which one is responsible. Instrument segmentation solves a different part: it tells us where each tool is, but not what it is doing. The rightmost panel joins these two kinds of information. Each interaction is attached to the mask of the responsible instrument instance. That is the task studied throughout this thesis. The target is a semantic anatomical label; this does not mean we have segmented all the tissue.',
      ['chapters/01_introduction.tex','figs/triplet_grounding/instance_recognition_triplet_segmentation.pdf','IPCAI deck, slides 3–4'])
figure('figs/triplet_grounding/instance_recognition_triplet_segmentation.pdf',.7,1.8,11.95,3.95)
takeaway('Grounding connects “what is happening?” to “which visible instrument is responsible?”');end()

start('The review maps how surgical vision tasks connect','Contribution 1 · multitask learning review',65,
      'The PhD began with a broad review of multitask learning in minimally invasive surgical vision. I organised work across perceptual tasks, tracking and camera control, workflow and activity understanding, and post-surgical analysis. These tasks need complementary information: objects, motion, procedural context, and interactions. Multitask learning can learn related outputs together, commonly through a shared encoder and separate heads. However, simply combining tasks does not guarantee improvements. The review identified negative transfer as well as benefits, and found that surgical work often used hard parameter sharing and simple weighted losses. Its contribution is both a map of applications and a critical discussion of how task relationships, optimisation, and annotation availability shape progress.',
      ['chapters/papers/02_review_paper.tex','TPC presentation.pptx, slides 5–9'])
card(.6,1.85,5.95,1.8,'Perception','Instruments, anatomy, depth and motion')
card(6.75,1.85,5.95,1.8,'Tracking and control','Pose, gaze and camera guidance')
card(.6,3.9,5.95,1.8,'Workflow and activity','Phases, steps and action triplets')
card(6.75,3.9,5.95,1.8,'Post-surgical analysis','Skill assessment and report generation')
takeaway('Related outputs can help each other; task compatibility and supervision still matter.');end()

start('The review exposes a missing spatial connection','From literature review to the thesis research gap',55,
      'The key transition from the review to the rest of the thesis is the gap between semantic richness and spatial evidence. Workflow labels say where we are in an operation. Triplets describe finer actions and anatomical targets. But frame-level triplets still leave the responsible instrument ambiguous. At the same time, large-scale instrument instance annotations for routine laparoscopic surgery were scarce. This points to two connected requirements: distinguish individual instruments, then align the action labels with those instances. Multitask learning remains useful background, but grounding becomes the organising problem of the thesis. The dataset work is therefore essential supervision for the later interaction models, rather than a separate side project.',
      ['chapters/papers/02_review_paper.tex, conclusion','chapters/01_introduction.tex, research gap'])
card(.6,1.85,3.9,2.35,'Frame-level meaning','Triplets describe the action and its anatomical context.')
card(4.7,1.85,3.9,2.35,'Instance evidence','Masks distinguish individual tools, including repeated classes.')
card(8.8,1.85,3.9,2.35,'Missing association','Each verb–target pair must be linked to the correct tool.')
text(.8,4.75,11.8,1.05,'Research direction: build the spatial supervision, align the interactions, then model the grounded relationship.',26,NAVY,True)
takeaway('Semantic labels without masks leave the actor ambiguous; masks without actions leave its role unknown.');end()

start('CholecInstanceSeg: distinguish individual instruments','Contribution 2 · instrument instance segmentation dataset',65,
      'CholecInstanceSeg establishes the spatial foundation. It contains 41,933 annotated frames from 85 procedures, with 64,483 instrument instances. The source videos come from the existing Cholec dataset family, and the annotations add semantic classes and instance identities. To scale the work, annotation is iterative: initialise masks, train a model, generate proposals, review and correct them, and retrain. Human review and quality control remain part of this process. The important advance for the thesis is that repeated or overlapping tools can be distinguished. Dataset limitations include class imbalance and difficult visual conditions. These masks enable the next step: connecting an instrument to its surgical role.',
      ['Abstract.tex','chapters/papers/cholecinstanceseg/methods.tex','chapters/papers/cholecinstanceseg/technical_validation.tex','poster_presentation.pptx'])
for x,num,label in [( .7,'41,933','annotated frames'),(4.9,'85','procedures'),(9.1,'64,483','instrument instances')]:
    rect(x,1.85,3.5,1.35,LIGHT);text(x+.2,2.0,3.1,.55,num,34,TEAL,True);text(x+.2,2.67,3.1,.35,label,17)
figure('figs/cholecinstanceseg/train-retrain-annotate.pdf',.7,3.55,11.95,2.0)
takeaway('Instance masks provide the spatial foundation for grounded action understanding.');end()

start('CholecTriplet-Seg: link masks to surgical interactions','Contribution 3 · dataset and triplet-segmentation benchmark',65,
      'CholecTriplet-Seg aligns CholecInstanceSeg masks with CholecT50 triplet labels, using manual alignment and quality control. It includes 30,955 frames, 49,866 grounded triplets, and 50 videos. Alignment is particularly important when multiple instruments share a class, or when mask labels and frame-level labels disagree. The prediction target is an instrument mask with its instrument, verb, and target labels. Evaluation checks the spatial match and the semantic labels together. Complete-triplet average precision is therefore demanding: the localisation and all three components must be correct for the same instance. Component scores help diagnose the errors, but they are not substitutes for complete interaction performance.',
      ['Abstract.tex','chapters/papers/targetfusionnet/main_body.tex','IPCAI deck, slides 4–5'])
text(.7,1.85,12,1.05,'CholecInstanceSeg masks + CholecT50 triplets\n+ manual alignment and quality control',26,NAVY,True)
for x,num,label in [(.7,'30,955','annotated frames'),(4.9,'49,866','grounded triplets'),(9.1,'50','videos')]:
    rect(x,3.1,3.5,1.35,LIGHT);text(x+.2,3.25,3.1,.55,num,34,TEAL,True);text(x+.2,3.91,3.1,.35,label,17)
text(.8,4.9,11.8,.85,'One prediction = instrument mask + instrument class + verb + anatomical target',25,NAVY,True)
takeaway('Complete-triplet AP requires the spatial match and all three labels to be correct together.');end()

start('TargetFusionNet adds weak anatomical context','Contribution 4 · target-aware grounded prediction',65,
      'Anatomical target prediction is a principal bottleneck. Tissue is deformable, partially visible, and often difficult to distinguish from local appearance alone. TargetFusionNet extends Mask2Former with a second source of information: coarse anatomy predictions from EndoViT. Gated fusion allows instrument queries to use these predictions as contextual cues. They are weak priors, not reference anatomy annotations, and their categories do not exactly match our target vocabulary. Under mask-based complete-triplet evaluation, the strongly supervised Mask2Former-Triplet baseline reaches 12.23 AP; TargetFusionNet reaches 13.47. The broad lesson is that strong instance supervision is foundational, and anatomical context provides an additional gain in this architecture.',
      ['chapters/papers/targetfusionnet/main_body.tex','Abstract.tex','IPCAI deck, slides 7–8'])
figure('figs/triplet_grounding/target_fusion.pdf',.65,1.8,7.6,4.25)
card(8.5,1.95,4.15,2.35,'Target-aware fusion','Instrument queries combine image features with coarse anatomy cues.')
text(8.7,4.65,3.8,.55,'12.23 → 13.47',31,TEAL,True)
text(8.7,5.27,3.8,.6,'Complete-triplet mask AP\nBaseline → TargetFusionNet',17)
takeaway('Strong instance supervision enables grounding; gated anatomy cues improve the evaluated baseline.');end()

start('Two stages separate localisation from interaction prediction','Contribution 5 · multimodal interaction prediction',65,
      'The final chapter asks whether localisation and interaction prediction can be specialised into separate stages. Stage one uses Mask2Former to segment and classify instruments. The image is then overlaid with the masks and instrument identifiers. Stage two uses Qwen3-VL-8B-Instruct, adapted with supervised fine-tuning, to predict a verb and a target for each supplied instance. The instrument identity is fixed by stage one and cannot be corrected by stage two. The output is structured and linked by instance IDs. I also examine sequential dialogue, short temporal context, anatomical prompts, ontology guidance, and distilled rationale supervision. These alternatives test whether a different task formulation or source of context improves the semantic prediction.',
      ['chapters/papers/two_stage/main_body.tex','phd_papers/sources/Two_Stage_Surgical_Triplet_Segmentation/figures/two_stage_pipeline.pdf'])
figure('phd_papers/sources/Two_Stage_Surgical_Triplet_Segmentation/figures/two_stage_pipeline.pdf',.65,1.75,7.0,4.5)
card(7.95,1.95,4.7,1.9,'Stage 1: Mask2Former','Predict instrument masks and classes.')
card(7.95,4.08,4.7,1.9,'Stage 2: Qwen3-VL','Predict one verb–target pair per supplied instance.')
takeaway('The complete triplet combines the Stage 1 instrument with the Stage 2 verb and target.');end()

start('Better components do not guarantee better complete triplets','Results · keep the two evaluation settings separate',80,
      'There are two evaluation settings, and it is important not to mix them. With ground-truth instruments supplied, direct supervised fine-tuning achieves 14.56 complete-triplet AP. Sequential dialogue gives 15.31, and three-frame context gives the highest observed score, 15.89. These are exploratory observations, not established significant improvements, and these adaptations were not evaluated in the complete predicted-instance pipeline. With predicted instruments, direct fine-tuning reaches 13.52 AP, essentially comparable to TargetFusionNet at 13.47. Verb and target component scores are higher, as shown in the table, but the complete relationship still needs every element to be correct together. Five-frame context and rationale-only supervision do not improve the result in the tested formulations. More context or longer generated explanations therefore does not automatically help.',
      ['chapters/papers/two_stage/main_body.tex, principal-adaptations and predicted-grounding tables'])
text(.7,1.7,12,.4,'COMPLETE SYSTEM · predicted instrument instances · mask-based AP',16,TEAL,True)
headers=['Method','Verb AP','Target AP','Triplet AP'];xs=[.9,7.15,9.05,11.0];ws=[6.0,1.6,1.65,1.6]
for row,values in enumerate([headers,['Mask2Former-Triplet','45.61','20.75','12.23'],['TargetFusionNet','46.27','21.55','13.47'],['Two-stage Direct SFT','51.89','25.27','13.52']]):
    yy=2.25+row*.58;rect(.7,yy,11.95,.55,LIGHT if row%2==0 else 'FFFFFF')
    for x,w,value in zip(xs,ws,values):text(x,yy+.08,w,.39,value,19,TEAL if row==0 else INK,row in [0,3])
text(.8,4.75,12,.4,'STAGE 2 ONLY · ground-truth instrument instances supplied',16,TEAL,True)
text(.8,5.35,11.8,.7,'Direct SFT: 14.56     |     Sequential dialogue: 15.31     |     3-frame context: 15.89',21)
takeaway('13.52 vs 13.47 is comparable complete-triplet AP; component gains are a separate finding.');end()

start('Grounding makes surgical interactions spatially explicit','Conclusion · contributions and next steps',50,
      'The take-home message is that recognising an interaction is not the same as identifying its actor. This thesis connects the two by providing instrument instance annotations, aligning them with triplets, and developing both target-aware and two-stage models. The review motivates this progression and gives the broader surgical-vision context. The main remaining challenge is robust target prediction and consistent complete-triplet assignment. Limitations include class imbalance, visual ambiguity, dependence on first-stage masks, and evaluation restricted to laparoscopic cholecystectomy. Next steps are temporal instance association, evaluation of the promising adaptations with predicted instruments, and testing across other procedures. The work provides a research foundation for interpretable surgical scene understanding; downstream clinical applications still need their own validation.',
      ['chapters/06_discussion_conclusion.tex','chapters/papers/two_stage/main_body.tex, limitations'])
card(.65,1.9,6.1,3.85,'What this thesis contributes','A review that identifies the grounding gap\nInstrument-instance and grounded-triplet datasets\nTarget-aware and two-stage interaction models')
card(7.0,1.9,5.65,3.85,'What comes next','More reliable target prediction\nTemporal instance association\nTesting across procedures and settings')
takeaway('Identify each instrument, then connect it to what it does and what it acts on.');end()

assert len(prs.slides)==10
assert sum(n[2] for n in notes)==600
prs.save(OUT/'Thesis_overview_10min.pptx');pdf.save()
elapsed=0;parts=['# Thesis overview: 10-minute speaker notes','Exactly 10 slides; planned speaking time: 10:00.','The thesis sources are authoritative where older decks differ.']
for idx,title,seconds,script in notes:
    parts.append(f'\n## {idx}. {title}\n\n{elapsed//60}:{elapsed%60:02d}–{(elapsed+seconds)//60}:{(elapsed+seconds)%60:02d} ({seconds} seconds)\n\n{script}')
    elapsed+=seconds
(OUT/'Thesis_overview_10min_notes.md').write_text('\n\n'.join(parts),encoding='utf-8')
(OUT/'Thesis_overview_10min_sources.json').write_text(json.dumps(provenance,indent=2,ensure_ascii=False),encoding='utf-8')
print('Created editable PowerPoint, matching 10-page PDF, speaker notes and source map.')
